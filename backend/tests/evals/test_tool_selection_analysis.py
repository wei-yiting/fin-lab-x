"""Tests for the tool-selection analysis over recorded agent runs.

Label fixtures are rows 1–5 as decided in the annotation pass (rules R1–R4);
row 4 is the coverage-gap row (TSMC / ASML file 20-F, the SEC tools read
10-K only).
"""

from __future__ import annotations

import csv
import io
import json
import tarfile
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from backend.evals.experiments.tool_selection.analyze import (
    LabelError,
    load_labels,
    load_records,
    main,
    normalize_ticker,
    parse_tickers,
    score_run,
    summarize,
)

HEADER = [
    "id",
    "題目",
    "公司",
    "預期 tickers（請確認）",
    "出題時的 best source",
    "類別",
    "Finnhub 報價",
    "Finnhub 財務指標",
    "Tavily 新聞",
    "SEC 10-K",
    "備註",
    "Coverage gap",
]
DECIDED = {
    1: ("TSLA", "可接受", "可接受", "必要", "必要", "否"),
    2: ("DIS", "不必要", "不必要", "必要", "必要", "否"),
    3: ("BA", "不必要", "不必要", "必要", "必要", "否"),
    4: ("TSM (20-F), ASML (20-F)", "不必要", "不必要", "必要", "必要", "是"),
    5: ("INTC", "不必要", "擇一必要", "必要", "擇一必要", "否"),
}

QUOTE = "finnhub_stock_quote"
FIN = "finnhub_company_basic_financials"
NEWS = "tavily_financial_search"
LIST = "sec_filing_list_sections"
SECTION = "sec_filing_get_section"
SEARCH = "sec_filing_search"


def write_sheet(path: Path, rows: dict[int, tuple[Any, ...]], total: int = 7) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "說明"
    ws = wb.create_sheet("標註")
    ws.append(HEADER)
    for row_id in range(1, total + 1):
        tickers, *labels, gap = rows.get(row_id, ("XYZ", None, None, None, None, None))
        ws.append([row_id, "q", "c", tickers, "mixed", "cat", *labels, None, gap])
    wb.save(path)
    return path


@pytest.fixture
def labels_path(tmp_path: Path) -> Path:
    rows = dict(DECIDED)
    rows[7] = ("JPM", "必要", None, None, None, None)
    return write_sheet(tmp_path / "labels.xlsx", rows)


def call(tool: str, ticker: str | None = None, error: str | None = None, **args: Any):
    if ticker is not None:
        args["ticker"] = ticker
    out: dict[str, Any] = {"tool": tool, "args": args}
    if error is None:
        out["result"] = "{}"
    else:
        out["error"] = error
    return out


def record(
    row_id: str, arm: str, rep: int, calls: list[dict[str, Any]], response: str = ""
) -> dict[str, Any]:
    return {
        "row_id": row_id,
        "arm": arm,
        "repetition": rep,
        "output": {"response": response, "tool_outputs": calls},
    }


# --- labels -----------------------------------------------------------------


def test_load_labels_keeps_only_fully_labelled_rows(labels_path: Path) -> None:
    sheet = load_labels(labels_path)

    assert sorted(sheet.rows, key=int) == ["1", "2", "3", "4", "5"]
    assert sheet.total_rows == 7
    assert sheet.incomplete == ["7"]
    assert sheet.gap_column_found


def test_load_labels_reads_sources_tickers_and_gap_flag(labels_path: Path) -> None:
    rows = load_labels(labels_path).rows

    assert rows["5"].labels == {
        "quote": "不必要",
        "financials": "擇一必要",
        "news": "必要",
        "sec10k": "擇一必要",
    }
    assert rows["4"].tickers == {"TSM", "ASML"}
    assert rows["4"].coverage_gap
    assert not rows["1"].coverage_gap


def test_load_labels_rejects_unknown_label_value(tmp_path: Path) -> None:
    path = write_sheet(
        tmp_path / "bad.xlsx", {1: ("TSLA", "可接受", "必須", "必要", "必要", "否")}
    )

    with pytest.raises(LabelError, match="row 1"):
        load_labels(path)


def test_load_labels_rejects_single_either_required_source(tmp_path: Path) -> None:
    path = write_sheet(
        tmp_path / "bad.xlsx", {1: ("TSLA", "可接受", "擇一必要", "必要", "必要", "否")}
    )

    with pytest.raises(LabelError, match="擇一必要"):
        load_labels(path)


def test_load_labels_without_gap_column_marks_no_gap_rows(tmp_path: Path) -> None:
    path = write_sheet(tmp_path / "labels.xlsx", DECIDED)
    wb = openpyxl.load_workbook(path)
    wb["標註"].delete_cols(len(HEADER))
    wb.save(path)

    sheet = load_labels(path)

    assert not sheet.gap_column_found
    assert not any(r.coverage_gap for r in sheet.rows.values())


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("TSLA", {"TSLA"}),
        ("TSM (20-F), ASML (20-F)", {"TSM", "ASML"}),
        ("BA, AIR.PA (非 SEC)", {"BA", "AIR.PA"}),
        ("AAPL、MSFT, googl", {"AAPL", "MSFT", "GOOGL"}),
    ],
)
def test_parse_tickers(raw: str, expected: set[str]) -> None:
    assert parse_tickers(raw) == expected


@pytest.mark.parametrize(
    ("a", "b"),
    [("GOOG", "GOOGL"), ("BRK-B", "BRK.B"), ("brk.b", "BRK.B"), ("BRK.A", "BRK.B")],
)
def test_normalize_ticker_treats_aliases_as_equal(a: str, b: str) -> None:
    assert normalize_ticker(a) == normalize_ticker(b)


def test_normalize_ticker_keeps_distinct_issuers_apart() -> None:
    assert normalize_ticker("AIR") != normalize_ticker("AIR.PA")


# --- per-run scoring --------------------------------------------------------


def test_recall_needs_every_required_and_one_either_source(labels_path: Path) -> None:
    row5 = load_labels(labels_path).rows["5"]

    with_fin = score_run(
        record("5", "a", 1, [call(NEWS, "INTC"), call(FIN, "INTC")]), row5
    )
    with_sec = score_run(
        record("5", "a", 1, [call(NEWS, "INTC"), call(SECTION, "INTC")]), row5
    )
    news_only = score_run(record("5", "a", 1, [call(NEWS, "INTC")]), row5)
    no_news = score_run(record("5", "a", 1, [call(FIN, "INTC")]), row5)

    assert with_fin["recall"] is True
    assert with_sec["recall"] is True
    assert news_only["recall"] is False
    assert no_news["recall"] is False


def test_any_sec_tool_counts_as_the_10k_source(labels_path: Path) -> None:
    row2 = load_labels(labels_path).rows["2"]

    for sec_tool in (LIST, SECTION, SEARCH):
        run = score_run(
            record("2", "a", 1, [call(NEWS, "DIS"), call(sec_tool, "DIS")]), row2
        )
        assert run["recall"] is True
        assert run["used_sec10k"] is True


def test_unnecessary_flags_and_counts_calls_to_unnecessary_sources(
    labels_path: Path,
) -> None:
    rows = load_labels(labels_path).rows
    calls = [
        call(QUOTE, "INTC"),
        call(QUOTE, "INTC"),
        call(NEWS, "INTC"),
        call(FIN, "INTC"),
    ]

    row5_run = score_run(record("5", "a", 1, calls), rows["5"])
    row1_run = score_run(record("1", "a", 1, calls), rows["1"])

    assert row5_run["unnecessary"] is True
    assert row5_run["unnecessary_calls"] == 2
    assert row1_run["unnecessary"] is False
    assert row1_run["unnecessary_calls"] == 0


def test_errored_only_source_counts_as_used_but_not_for_success_only_recall(
    labels_path: Path,
) -> None:
    row5 = load_labels(labels_path).rows["5"]
    calls = [call(NEWS, "INTC"), call(SECTION, "INTC", error="Section '7' not found")]

    run = score_run(record("5", "a", 1, calls), row5)

    assert run["recall"] is True
    assert run["recall_success_only"] is False
    assert run["errored_only_sources"] == "sec10k"


def test_ticker_accuracy_over_calls_that_pass_a_ticker(labels_path: Path) -> None:
    row4 = load_labels(labels_path).rows["4"]
    calls = [
        call(NEWS, "tsm"),
        call(NEWS, "ASML"),
        call(NEWS, "NVDA"),
        call(NEWS, query="AI export controls"),
    ]

    run = score_run(record("4", "a", 1, calls), row4)

    assert run["ticker_calls"] == 3
    assert run["ticker_correct"] == 2
    assert run["ticker_accuracy"] == pytest.approx(2 / 3)
    assert run["off_list_tickers"] == "NVDA"


def test_insufficient_evidence_heuristic(labels_path: Path) -> None:
    row4 = load_labels(labels_path).rows["4"]

    stated = score_run(
        record("4", "a", 1, [], response="目前證據不足以量化兩家公司的曝險。"), row4
    )
    silent = score_run(record("4", "a", 1, [], response="台積電暴露較直接。"), row4)

    assert stated["insufficient_evidence_stated"] is True
    assert "證據不足" in stated["insufficient_evidence_match"]
    assert silent["insufficient_evidence_stated"] is False


@pytest.mark.parametrize(
    ("response", "stated"),
    [
        ("現有資料不足以證明它比艾司摩爾更嚴重", True),
        ("SEC 工具無法取得其 10-K", True),
        ("但我沒有取得足夠的台積電中國營收", True),
        ("### 三、數據品質與資訊限制說明", True),
        ("*資料品質與限制說明：兩家公司均申報 20-F*", True),
        ("美荷持續擴大出口限制，ASML 中國營收下滑", False),
        ("台積電受限於美方規定無法為特定中國客戶代工", False),
    ],
)
def test_insufficient_evidence_phrases(
    labels_path: Path, response: str, stated: bool
) -> None:
    row4 = load_labels(labels_path).rows["4"]

    run = score_run(record("4", "a", 1, [], response=response), row4)

    assert run["insufficient_evidence_stated"] is stated


def test_repeated_calls_after_a_failure(labels_path: Path) -> None:
    row4 = load_labels(labels_path).rows["4"]
    fpi = "Ticker TSM appears to be a foreign private issuer that files 20-F"
    calls = [
        call(LIST, "TSM", error=fpi, doc_type="10-K"),
        call(LIST, "TSM", error=fpi, doc_type="10-K"),
        call(LIST, "TSM", doc_type="20-F", error=fpi),
        call(SEARCH, "TSM", error=fpi, query="export"),
        call(LIST, "ASML", doc_type="10-K"),
    ]

    run = score_run(record("4", "a", 1, calls), row4)

    assert run["repeated_failed_calls"] == 1
    assert run["retries_after_source_failure"] == 3


# --- aggregation ------------------------------------------------------------


def test_summarize_averages_reps_per_row_then_excludes_gap_rows(
    labels_path: Path,
) -> None:
    sheet = load_labels(labels_path)
    good = [call(NEWS, "X"), call(SECTION, "X")]
    bad = [call(NEWS, "X")]
    records = [
        record("2", "cfg", 1, good),
        record("2", "cfg", 2, good),
        record("2", "cfg", 3, bad),
        record("3", "cfg", 1, good),
        record("4", "cfg", 1, bad, response="證據不足"),
        record("6", "cfg", 1, bad),
    ]
    runs = [
        score_run(r, sheet.rows.get(r["row_id"])) | {"config": "cfg"} for r in records
    ]

    summary = summarize(runs, sheet)
    cfg = summary["configs"]["cfg"]

    assert cfg["routing_rows"] == 2
    point, lo, hi = cfg["recall"]
    assert point == pytest.approx((2 / 3 + 1) / 2)
    assert lo <= point <= hi
    assert cfg["gap_rows"] == 1
    assert cfg["insufficient_evidence_stated"][0] == 1.0
    assert cfg["median_tool_calls"] == 1.5
    assert cfg["runs"] == 6


def test_score_run_leaves_label_metrics_blank_for_unlabelled_rows() -> None:
    run = score_run(record("6", "cfg", 1, [call(NEWS, "X")]), None)

    assert run["labelled"] is False
    assert run["recall"] == ""
    assert run["tool_calls_total"] == 1


# --- archives + CLI -----------------------------------------------------------


def write_archive(archive_dir: Path, records: list[dict[str, Any]]) -> Path:
    """Mimics the real archives, which macOS tar wrote with an AppleDouble
    ``._<name>`` member (binary xattr data) beside every file."""
    archive_dir.mkdir(parents=True)

    def add(tar: tarfile.TarFile, name: str, data: bytes) -> None:
        info = tarfile.TarInfo(name)
        info.size = len(data)
        tar.addfile(info, io.BytesIO(data))

    with tarfile.open(archive_dir / "raw.tar.gz", "w:gz") as tar:
        for r in records:
            folder = f"raw/rep{r['repetition']}"
            name = f"row{int(r['row_id']):02d}__{r['arm']}.json"
            add(tar, f"{folder}/._{name}", b"\x00\x05\x16\x07\x00\x02Mac OS X")
            add(tar, f"{folder}/{name}", json.dumps(r).encode())
    return archive_dir


def test_load_records_qualifies_arm_names_shared_across_archives(
    tmp_path: Path,
) -> None:
    first = write_archive(
        tmp_path / "2026-10-22-model-reasoning-benchmark",
        [record("1", "luna_medium", 1, []), record("1", "luna_none", 1, [])],
    )
    second = write_archive(
        tmp_path / "2026-10-22-rag-bundle-ablation",
        [record("1", "luna_medium", 1, []), record("1", "reader_luna_medium", 1, [])],
    )

    configs = {r["config"] for r in load_records([first, second])}

    assert configs == {
        "luna_medium @ model-reasoning-benchmark",
        "luna_none",
        "luna_medium @ rag-bundle-ablation",
        "reader_luna_medium",
    }


def test_main_writes_outputs_and_reports_scored_rows(
    tmp_path: Path, labels_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    archive = write_archive(
        tmp_path / "2026-10-22-exp",
        [
            record("1", "cfg", 1, [call(NEWS, "TSLA"), call(SECTION, "TSLA")]),
            record("4", "cfg", 1, [call(LIST, "TSM", error="20-F")], "證據不足"),
            record("6", "cfg", 1, [call(NEWS, "XYZ")]),
        ],
    )
    out_dir = tmp_path / "out"

    assert (
        main(
            [
                "--labels",
                str(labels_path),
                "--archives",
                str(archive),
                "--out-dir",
                str(out_dir),
            ]
        )
        == 0
    )

    stdout = capsys.readouterr().out
    assert "5 of 7 rows labelled" in stdout
    assert "4 scored for routing" in stdout
    tables = (out_dir / "tables.md").read_text()
    assert "PRELIMINARY" in tables
    assert "| 4 | cfg | 1 | 證據不足 |" in tables
    with open(out_dir / "metrics.csv", newline="") as f:
        assert len(list(csv.DictReader(f))) == 3
